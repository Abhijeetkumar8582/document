"""Low-level readers: text layer, tables, page rendering, DOCX and plain text."""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from ..config import settings


@dataclass
class Extraction:
    text: str
    page_count: int = 0
    method: str = "text"
    tables: list[list[list[str | None]]] = field(default_factory=list)
    form_fields: dict[str, str] = field(default_factory=dict)
    structured: dict | None = None  # fields from an LLM pass, if any
    structured_subjects: list[dict] | None = None


class UnsupportedFile(Exception):
    pass


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
MIME = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
    ".bmp": "image/bmp",
}


def detect_kind(filename: str, content_type: str | None) -> str:
    ext = Path(filename).suffix.lower()
    ct = (content_type or "").lower()
    if ext == ".pdf" or ct == "application/pdf":
        return "pdf"
    if ext in IMAGE_EXTS or ct.startswith("image/"):
        return "image"
    if ext == ".docx" or ct == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return "docx"
    if ext in {".txt", ".csv", ".md"} or ct.startswith("text/"):
        return "text"
    raise UnsupportedFile(f"Unsupported file type: {ext or ct or 'unknown'}")


def mime_for(filename: str, kind: str) -> str:
    return MIME.get(Path(filename).suffix.lower(), "image/png" if kind == "image" else "application/pdf")


# --- PDF -----------------------------------------------------------------------


def pdf_text_layer(data: bytes) -> Extraction:
    """Text and tables from the PDF's own text layer. Empty text means it is a scan."""
    import pdfplumber

    texts: list[str] = []
    tables: list[list[list[str | None]]] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            texts.append(page.extract_text(x_tolerance=1.5, y_tolerance=3) or "")
            try:
                for table in page.extract_tables():
                    if table and len(table) > 1:
                        tables.append(table)
            except Exception:
                pass
    text = "\n".join(texts).strip()
    if len(text) < settings.min_text_chars_per_page * max(page_count, 1):
        try:
            from pypdf import PdfReader

            alt = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages).strip()
            if len(alt) > len(text):
                text = alt
        except Exception:
            pass
    return Extraction(text=text, page_count=page_count, method="text", tables=tables)


def render_pages(data: bytes, kind: str, max_pages: int = 20) -> list:
    """PIL images for each page (PDF) or the single image."""
    from PIL import Image

    if kind == "image":
        return [Image.open(io.BytesIO(data)).convert("RGB")]
    import pdfplumber

    images = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages[:max_pages]:
            images.append(page.to_image(resolution=settings.render_dpi).original.convert("RGB"))
    return images


def has_text_layer(extraction: Extraction) -> bool:
    return len(extraction.text) >= settings.min_text_chars_per_page * max(extraction.page_count, 1)


# --- Local OCR (last resort) ------------------------------------------------------


def tesseract_available() -> bool:
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def ocr_images(images: list) -> str:
    import pytesseract

    return "\n".join(pytesseract.image_to_string(img) for img in images).strip()


# --- DOCX / text --------------------------------------------------------------------


def docx_text(data: bytes) -> Extraction:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    tables: list[list[list[str | None]]] = []
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        tables.append(rows)
        for row in rows:
            lines.append("  ".join(row))
    return Extraction(text="\n".join(lines).strip(), page_count=1, method="text", tables=tables)


def plain_text(data: bytes) -> Extraction:
    return Extraction(text=data.decode("utf-8", errors="replace"), page_count=1, method="text")
